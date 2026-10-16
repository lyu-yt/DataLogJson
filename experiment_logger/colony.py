"""Mouse profiles and read-only XLSX lookup with explicit identity selection."""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
import re

from jsonschema import Draft7Validator, FormatChecker

from .io import LoggerError, Cancelled, encode, hash_file, now, read_json, signature, uid
from .layout import mouse_folder
from .models import RESOURCE, SCHEMA_VERSION

COLONY_FIELDS = {"Mouse ID": "colony_id", "Strain": "strain", "Sex": "sex", "DOB": "date_of_birth",
                 "Genotype": "genotype", "Cage ID": "cage_id", "Use": "use_status",
                 "Labels": "labels", "Notes": "colony_notes", "SystemId": "system_id"}


def new_profile(mouse_id, author):
    mouse_folder(mouse_id)
    if not author or not author.strip():
        raise LoggerError("An operator name is required")
    stamp = now()
    return {"schema_version": SCHEMA_VERSION, "profile_type": "mouse", "profile_id": uid(),
            "mouse_id": mouse_id, "revision": 0, "created_at": stamp, "created_by": author,
            "updated_at": stamp, "updated_by": author,
            **{field: None for field in COLONY_FIELDS.values()},
            "notes": None, "surgeries": [], "import_history": []}


def require_profile(profile, schema=None):
    try:
        encode(profile)
    except (TypeError, ValueError) as exc:
        raise LoggerError(f"profile JSON: {exc}") from exc
    schema = schema or read_json(RESOURCE / "schemas" / "record.json")
    definition = schema.get("definitions", {}).get("mouse_profile")
    if not definition:
        raise LoggerError("The archived schema does not define mouse profiles")
    errors = [f"{'.'.join(map(str, e.absolute_path)) or 'profile'}: {e.message}"
              for e in Draft7Validator(definition, format_checker=FormatChecker()).iter_errors(profile)]
    if errors:
        raise LoggerError("\n".join(errors))
    surgeries = profile.get("surgeries", [])
    if len({s.get("surgery_id") for s in surgeries}) != len(surgeries):
        errors.append("surgeries: duplicate surgery IDs")
    history = profile["import_history"]
    if len({entry["import_id"] for entry in history}) != len(history):
        errors.append("import_history: duplicate import IDs")
    for i, entry in enumerate(history):
        if entry["matched_mouse_id"] != profile["mouse_id"]:
            errors.append(f"import_history.{i}.matched_mouse_id: does not match local mouse ID")
    birth = profile.get("date_of_birth")
    if birth:
        for i, surgery in enumerate(surgeries):
            if surgery.get("date") and surgery["date"] < birth:
                errors.append(f"surgeries.{i}.date: precedes date of birth")
    if errors:
        raise LoggerError("\n".join(errors))


def _text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() or None


def _birth_date(value, cell):
    if value is None or not str(value).strip():
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise LoggerError(f"{cell}: unrecognized DOB {text!r}; use YYYY-MM-DD or MM/DD/YYYY")


def search_workbook(path, mouse_id, progress=None, cancel=None):
    """Return matching rows, never silently map a prefixed ID or choose a duplicate.

    The supplied colony export declares an A1-only worksheet dimension despite
    containing many rows. Reset read-only dimensions and find actual headers.
    No Excel formulas or macros are executed and the workbook is never saved.
    """
    mouse_folder(mouse_id)
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise LoggerError("XLSX support is missing. Update the datalogjson Conda environment using environment.yml.") from exc
    path = Path(path)
    before = signature(path)
    checksum = hash_file(path, progress, cancel)["sha256"]
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    candidates, headers_found = [], False
    try:
        for sheet in book:
            sheet.reset_dimensions()
            header = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if cancel and cancel():
                    raise Cancelled("Colony search cancelled")
                names = {str(v).strip(): i for i, v in enumerate(row) if v is not None and str(v).strip()}
                if "Mouse ID" in names and "Strain" in names:
                    header = names
                    headers_found = True
                    continue
                if header is None:
                    continue
                def value(name):
                    i = header.get(name)
                    return row[i] if i is not None and i < len(row) else None
                colony_id = _text(value("Mouse ID"))
                match = re.fullmatch(r"([A-Za-z]*)([0-9]+)", colony_id or "")
                if not match or int(match.group(2)) != int(mouse_id):
                    continue
                fields, cells, original = {}, {}, {}
                for name, field in COLONY_FIELDS.items():
                    if name not in header:
                        continue
                    cell = f"{get_column_letter(header[name] + 1)}{row_number}"
                    v = value(name)
                    fields[field] = _birth_date(v, f"{sheet.title}!{cell}") if field == "date_of_birth" else _text(v)
                    cells[field] = cell
                    original[name] = v.isoformat() if isinstance(v, (date, datetime)) else str(v) if v is not None else None
                candidates.append({"mouse_id": mouse_id, "colony_id": colony_id, "fields": fields,
                    "has_prefix": bool(match.group(1)), "requires_confirmation": colony_id != mouse_id,
                    "source": {"workbook_filename": path.name, "workbook_sha256": checksum,
                               "sheet": sheet.title, "row": row_number, "cells": cells,
                               "source_values": original}, "read_at": now()})
    finally:
        book.close()
    if before != signature(path):
        raise LoggerError("Colony workbook changed during lookup; retry with the completed file")
    if not headers_found:
        raise LoggerError("No worksheet contains the required Mouse ID and Strain column headers")
    if len(candidates) > 1:
        for candidate in candidates:
            candidate["requires_confirmation"] = True
    return candidates


def import_candidate(profile, candidate, author, *, confirmed=False):
    if not confirmed:
        raise LoggerError("Select and confirm the colony row before importing; prefixed or duplicate IDs are never chosen automatically")
    if candidate["mouse_id"] != profile["mouse_id"]:
        raise LoggerError("The selected colony row belongs to another local mouse ID")
    result = deepcopy(profile)
    for field in COLONY_FIELDS.values():
        if field in candidate["fields"]:
            result[field] = candidate["fields"][field]
    result["import_history"].append({"import_id": uid(), "imported_at": now(), "imported_by": author,
        "selected_colony_id": candidate["colony_id"], "matched_mouse_id": profile["mouse_id"],
        "confirmed_selection": True, "read_at": candidate["read_at"],
        "source": deepcopy(candidate["source"]), "imported_values": deepcopy(candidate["fields"])})
    require_profile(result)
    return result
